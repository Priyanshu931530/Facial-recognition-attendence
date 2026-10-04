import face_recognition
from tkinter import ttk, messagebox, filedialog
import cv2
import openpyxl
from datetime import datetime, timedelta
import os
import mysql.connector

def is_record_exists(sheet, name, current_date):
    for row in sheet.iter_rows(min_row=2, max_col=2):
        if row[0].value == name and row[1].value == current_date:
            return True
    return False

class FaceRecognition:
    def __init__(self, attendance_file='attendance.xlsx', webcam_index=0):
        self.known_faces = []
        self.known_names = []
        self.face_locations = []
        self.face_encodings = []
        self.face_names = []
        self.process_this_frame = True
        self.attendance_file = attendance_file
        self.webcam_index = webcam_index
        self.attendance_cooldown = {}

        try:
            self.conn = mysql.connector.connect(
                host="localhost",
                user="root",
                password="1234",
                database="student_management"
            )
            self.cursor = self.conn.cursor()
        except mysql.connector.Error as err:
            messagebox.showerror("Database Error", f"Error connecting to database: {err}")
            quit()

        self.load_known_faces_from_database()

        if not os.path.isfile(self.attendance_file):
            self.create_attendance_file()

        self.wb = openpyxl.load_workbook(self.attendance_file)
        self.sheet = self.wb.active

    def create_attendance_file(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append(["NAME", "DATE & TIME"])
        wb.save(self.attendance_file)

    def load_known_faces_from_database(self):
        try:
            self.cursor.execute("SELECT name, image_path FROM student_details")
            rows = self.cursor.fetchall()

            for name, image_path in rows:
                if os.path.exists(image_path):
                    try:
                        image = face_recognition.load_image_file(image_path)
                        encoding = face_recognition.face_encodings(image)[0]
                        self.known_faces.append(encoding)
                        self.known_names.append(name)
                    except Exception as e:
                        print(f"Error processing image {image_path}: {e}")
                else:
                    print(f"Image path does not exist: {image_path}")
        except mysql.connector.Error as err:
            print(f"Database query error: {err}")

    def mark_attendance(self, sheet, name):
        current_date = datetime.now().strftime("%Y-%m-%d %H:%M")
        if not is_record_exists(sheet, name, current_date):
            sheet.append([name, current_date])
            return True
        return False

    def start_face_detection(self):
        video_capture = cv2.VideoCapture(self.webcam_index)

        try:
            while True:
                ret, frame = video_capture.read()
                if not ret:
                    break

                small_frame = cv2.resize(frame, (0, 0), fx=0.25, fy=0.25)

                if self.process_this_frame:
                    self.detect_faces_and_mark_attendance(small_frame, self.sheet)

                self.process_this_frame = not self.process_this_frame

                self.display_results(frame)

                if cv2.waitKey(1) & 0xFF == ord('q'):
                    break

        except Exception as e:
            print(f"An error occurred: {str(e)}")

        finally:
            video_capture.release()
            cv2.destroyAllWindows()

    def detect_faces_and_mark_attendance(self, small_frame, sheet):
        self.face_locations = face_recognition.face_locations(small_frame)
        self.face_encodings = face_recognition.face_encodings(small_frame, self.face_locations)

        self.face_names = []
        for face_encoding in self.face_encodings:
            matches = face_recognition.compare_faces(self.known_faces, face_encoding)
            name = "Unknown"

            if True in matches:
                first_match_index = matches.index(True)
                name = self.known_names[first_match_index]

            self.face_names.append(name)

            if name != "Unknown" and self.should_mark_attendance(name):
                if self.mark_attendance(sheet, name):
                    self.wb.save(self.attendance_file)
                    self.attendance_cooldown[name] = datetime.now() + timedelta(minutes=1)

    def should_mark_attendance(self, name):
        if name in self.attendance_cooldown:
            return datetime.now() > self.attendance_cooldown[name]
        return True

    def display_results(self, frame):
        for (top, right, bottom, left), name in zip(self.face_locations, self.face_names):
            top *= 4
            right *= 4
            bottom *= 4
            left *= 4

            cv2.rectangle(frame, (left, top), (right, bottom), (0, 0, 255), 2)
            cv2.rectangle(frame, (left, bottom - 35), (right, bottom), (0, 0, 255), cv2.FILLED)

            font = cv2.FONT_HERSHEY_DUPLEX
            cv2.putText(frame, name, (left + 6, bottom - 6), font, 1.0, (255, 255, 255), 1)

        cv2.imshow('Video', frame)

if __name__ == "__main__":
    face_recognition_system = FaceRecognition()
    face_recognition_system.start_face_detection()
